"""Read user stories out of an uploaded file.

The text goes into the page's text box, so the user sees exactly what will be
sent and can edit it first. Nothing is sent to the model from here.

Spreadsheets need one judgement: the user's own workbook keeps stories and test
cases in the same file (sheets US1, TC1, US2, TC2). Sheets that look like a test
case table - a header with "TC ID" or "Expected Result" - are skipped, and the
reply says so.
"""

import io
import os
from typing import List, Tuple

SUPPORTED = (".txt", ".md", ".docx", ".xlsx", ".pdf")
MAX_BYTES = 5 * 1024 * 1024

_TEST_CASE_HEADERS = ("tc id", "test case id", "expected result", "expected output", "test id")


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    # Keep blank lines between paragraphs that were separated, drop runs of them.
    out, blank = [], False
    for line in lines:
        if line.strip():
            out.append(line.rstrip())
            blank = False
        elif not blank:
            out.append("")
            blank = True
    return "\n".join(out).strip()


def _xlsx(data: bytes) -> Tuple[str, List[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks, skipped = [], []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(cells)
        if not rows:
            continue
        header = " ".join(" ".join(r) for r in rows[:3]).lower()
        if any(h in header for h in _TEST_CASE_HEADERS):
            skipped.append(ws.title)
            continue
        blocks.append("\n".join(" ".join(r) for r in rows))
    wb.close()
    return "\n\n".join(blocks), skipped


def _pdf(data: bytes) -> str:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n\n".join((p.extract_text() or "").strip() for p in pdf.pages).strip()


def extract_text(filename: str, data: bytes) -> Tuple[str, List[str]]:
    """Return (text, notes). Raises ValueError with a plain message on failure."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in SUPPORTED:
        raise ValueError("%s isn't supported. Upload %s." % (ext or "That file type",
                                                             ", ".join(SUPPORTED)))
    if len(data) > MAX_BYTES:
        raise ValueError("%s is larger than 5 MB." % filename)
    notes: List[str] = []
    try:
        if ext in (".txt", ".md"):
            text = data.decode("utf-8", errors="replace").strip()
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".xlsx":
            text, skipped = _xlsx(data)
            if skipped:
                notes.append("Skipped sheet%s %s: %s look like test cases, not stories."
                             % ("s" if len(skipped) > 1 else "", ", ".join(skipped),
                                "they" if len(skipped) > 1 else "it"))
        else:
            text = _pdf(data)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Couldn't read %s (%s)." % (filename, type(exc).__name__))
    if not text:
        msg = "No text found in %s." % filename
        if ext == ".pdf":
            msg += " It may be a scanned image; export it as text or Word instead."
        raise ValueError(msg)
    return text, notes
