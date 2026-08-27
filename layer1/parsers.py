"""Layer 1 - document ingestion.

HLD 8.4 lists `IDocumentParser` as one of the extensibility seams: adding a
format means implementing an interface and registering it, never touching
extraction or orchestration.

Parsers produce `Block`s rather than one blob of text. A block is the natural
traceability unit - a user story, a spreadsheet row, a Jira issue - so every
Requirement can carry a `source_ref` that points back at something a human can
open (BR-02).

Failure is partial and visible, never silent (HLD 4.1): an unreadable page
becomes a warning on the ParsedDocument, not a missing requirement.
"""

import json
import os
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from typing import Dict, List, Optional

from layer1.contracts import Source


@dataclass
class Block:
    """One addressable chunk of a document."""

    ref: str
    text: str
    # A workflow the parser knows structurally (spreadsheet sheet name, Jira
    # component) but that does not appear in the block's own text.
    hint: Optional[str] = None


@dataclass
class ParsedDocument:
    name: str
    kind: str
    source: Source
    blocks: List[Block] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks)


# --- story splitting -------------------------------------------------------

# A line that starts a new story: a markdown heading, a bare ID, or a numbered
# item. Captures an ID (US-004, PROJ-1423, REQ 12) when one is present.
_ID = r"(?:US|STORY|REQ|AC|PROJ|TC)[-_ ]?\d+"
_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*)$")
_ID_LINE = re.compile(rf"^\s*[\-\*\d\.\)\s]*({_ID})\b[:\s\-—]*(.*)$", re.IGNORECASE)
_AS_A = re.compile(r"^\s*(?:\*\*)?as an?\b", re.IGNORECASE)
_ID_ANY = re.compile(_ID, re.IGNORECASE)


def split_story_blocks(text: str, fallback_prefix: str) -> List[Block]:
    """Split free text into story-sized blocks with a traceable ref each."""
    lines = text.splitlines()
    blocks: List[Block] = []
    current: List[str] = []
    current_ref: Optional[str] = None
    started_line = 1

    def flush() -> None:
        if not current:
            return
        body = "\n".join(current).strip()
        if body:
            ref = current_ref or f"{fallback_prefix}#L{started_line}"
            blocks.append(Block(ref=ref, text=body))

    for i, line in enumerate(lines, start=1):
        is_start = False
        ref_here: Optional[str] = None

        m_id = _ID_LINE.match(line)
        m_head = _HEADING.match(line)
        if m_id:
            is_start, ref_here = True, m_id.group(1).upper().replace("_", "-").replace(" ", "-")
        elif m_head:
            is_start = True
            found = _ID_ANY.search(m_head.group(1))
            ref_here = found.group(0).upper().replace("_", "-").replace(" ", "-") if found else None
        elif _AS_A.match(line) and any(_AS_A.match(c) for c in current):
            is_start = True

        if is_start and current:
            flush()
            current, current_ref, started_line = [], None, i

        if is_start and ref_here:
            current_ref = f"{fallback_prefix}#{ref_here}"
        if not current:
            started_line = i
        current.append(line)

    flush()

    # No structure at all - treat the whole document as one block so the LLM
    # path can still work, rather than returning nothing.
    if not blocks and text.strip():
        blocks = [Block(ref=f"{fallback_prefix}#whole", text=text.strip())]
    return blocks


# --- the seam --------------------------------------------------------------


class IDocumentParser(ABC):
    """Implement + register to support a new format. HLD 8.4."""

    extensions: List[str] = []
    kind: str = "unknown"
    source: Source = Source.DOCUMENT

    @abstractmethod
    def parse(self, path: str) -> ParsedDocument:
        ...

    def _doc(self, path: str) -> ParsedDocument:
        return ParsedDocument(name=os.path.basename(path), kind=self.kind, source=self.source)


class TextParser(IDocumentParser):
    extensions = [".md", ".txt", ".markdown"]
    kind = "markdown"

    def parse(self, path: str) -> ParsedDocument:
        doc = self._doc(path)
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            text = fh.read()
        doc.blocks = split_story_blocks(text, doc.name)
        return doc


class DocxParser(IDocumentParser):
    extensions = [".docx"]
    kind = "docx"

    def parse(self, path: str) -> ParsedDocument:
        import docx  # python-docx

        doc = self._doc(path)
        d = docx.Document(path)
        parts: List[str] = []
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            # Preserve heading structure so the splitter can see story starts.
            if style.startswith("Heading"):
                parts.append("## " + t)
            elif "List" in style:
                # Word stores the bullet as list formatting, not text. Restore
                # the marker so acceptance criteria are recognisable downstream.
                parts.append("- " + t)
            else:
                parts.append(t)
        doc.blocks = split_story_blocks("\n".join(parts), doc.name)

        # A table is a record set, like a spreadsheet: header row supplies the
        # labels, each remaining row becomes its own addressable block.
        for t_i, table in enumerate(d.tables, start=1):
            rows = [[c.text.strip() for c in row.cells] for row in table.rows]
            if len(rows) < 2:
                continue
            header = rows[0]
            for r_i, row in enumerate(rows[1:], start=2):
                pairs = [f"{h}: {v}" for h, v in zip(header, row) if h and v]
                if pairs:
                    doc.blocks.append(
                        Block(ref=f"{doc.name}#table{t_i}!row{r_i}", text="\n".join(pairs))
                    )
        return doc


class PdfParser(IDocumentParser):
    """Structure-aware PDF extraction.

    A flat `extract_text()` dump loses exactly the things Layer 1 keys on:
    headings that start a story, and bullets that mark acceptance criteria.
    (MarkItDown was evaluated for this and produced zero headings and zero
    bullets on a real Word-exported PDF, so it does not solve the problem.)

    PDFs carry no semantic markup, but they do carry the evidence needed to
    recover it: glyph font size, x-indent, and literal bullet characters. This
    parser rebuilds markdown structure from those, so a PDF reaches the rules
    engine looking like the markdown it was before it was printed.
    """

    extensions = [".pdf"]
    kind = "pdf"

    _BULLET_GLYPHS = "\u2022\u25aa\u25e6\u2023\u00b7\u2043-\u2013\u2014*"
    _NUMBERED = re.compile(r"^\s*\d{1,2}[\.\)]\s+\S")

    def parse(self, path: str) -> ParsedDocument:
        import pdfplumber

        doc = self._doc(path)
        lines: List[dict] = []
        with pdfplumber.open(path) as pdf:
            for n, page in enumerate(pdf.pages, start=1):
                try:
                    lines.extend(self._page_lines(page))
                except Exception as exc:  # noqa: BLE001 - partial failure stays visible
                    doc.warnings.append(f"page {n} unparsed: {exc}")

        if not lines:
            doc.warnings.append("no extractable text (scanned PDF? OCR not enabled)")
            doc.blocks = []
            return doc

        body_size = self._body_size(lines)
        doc.blocks = split_story_blocks(self._to_markdown(lines, body_size), doc.name)
        return doc

    def _page_lines(self, page) -> List[dict]:
        """Group words into lines, keeping font size and indent.

        Words rather than characters: joining raw chars reintroduces PDF
        kerning artefacts ("passwor d"), because a wide letter gap is
        indistinguishable from a space at the character level.
        """
        words = page.extract_words(extra_attrs=["size"], keep_blank_chars=False)
        rows: Dict[int, list] = {}
        for w in words:
            rows.setdefault(round(w["top"] / 2.0), []).append(w)

        out: List[dict] = []
        for key in sorted(rows):
            line = sorted(rows[key], key=lambda w: w["x0"])
            text = " ".join(w["text"] for w in line).strip()
            if not text:
                continue
            out.append({
                "text": text,
                "size": round(max(w.get("size", 0) or 0 for w in line), 1),
                "x0": round(min(w["x0"] for w in line)),
            })
        return out

    def _body_size(self, lines: List[dict]) -> float:
        """The most common font size is the body text; anything larger is a heading."""
        counts: Dict[float, int] = {}
        for ln in lines:
            counts[ln["size"]] = counts.get(ln["size"], 0) + len(ln["text"])
        return max(counts.items(), key=lambda kv: kv[1])[0] if counts else 11.0

    def _to_markdown(self, lines: List[dict], body_size: float) -> str:
        heading_size = body_size * 1.15
        base_indent = min((ln["x0"] for ln in lines), default=0)
        out: List[str] = []

        for ln in lines:
            text = ln["text"]
            stripped = text.lstrip(self._BULLET_GLYPHS + " \t")
            is_bullet = text[:1] in self._BULLET_GLYPHS and stripped
            is_numbered = bool(self._NUMBERED.match(text))

            if ln["size"] >= heading_size:
                out.append("## " + text)
            elif is_bullet:
                out.append("- " + stripped)
            elif is_numbered:
                out.append("- " + text)
            elif (
                out
                and ln["x0"] > base_indent + 8
                and out[-1].startswith("- ")
                and not out[-1].rstrip().endswith((".", ":", "!", "?"))
            ):
                # A wrapped continuation of the previous bullet, not a new line.
                out[-1] = out[-1].rstrip() + " " + text
            else:
                out.append(text)

        return "\n".join(out)


class XlsxParser(IDocumentParser):
    """Spreadsheet rows are already records - one row, one block."""

    extensions = [".xlsx", ".xlsm"]
    kind = "xlsx"

    def parse(self, path: str) -> ParsedDocument:
        from openpyxl import load_workbook

        doc = self._doc(path)
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header = [str(c).strip() if c is not None else "" for c in next(rows)]
            except StopIteration:
                continue
            for idx, row in enumerate(rows, start=2):
                pairs = [
                    f"{h}: {v}"
                    for h, v in zip(header, row)
                    if h and v is not None and str(v).strip()
                ]
                if not pairs:
                    continue
                doc.blocks.append(
                    Block(
                        ref=f"{doc.name}#{ws.title}!row{idx}",
                        text="\n".join(pairs),
                        hint=str(ws.title),
                    )
                )
        wb.close()
        return doc


class JiraJsonParser(IDocumentParser):
    """Jira export. Source is `jira`, and source_ref is the real issue key."""

    extensions = [".json"]
    kind = "jira"
    source = Source.JIRA

    def parse(self, path: str) -> ParsedDocument:
        doc = self._doc(path)
        with open(path, "r", encoding="utf-8") as fh:
            payload = json.load(fh)

        issues = payload.get("issues", payload) if isinstance(payload, dict) else payload
        if not isinstance(issues, list):
            doc.warnings.append("unrecognised JSON shape; expected a list of issues")
            return doc

        for i, issue in enumerate(issues, start=1):
            if not isinstance(issue, dict):
                doc.warnings.append(f"issue #{i}: expected an object, got {type(issue).__name__}")
                continue
            fields = issue.get("fields", issue)
            if not isinstance(fields, dict):
                doc.warnings.append(f"issue #{i}: 'fields' is not an object")
                continue
            key = issue.get("key") or fields.get("key") or f"ISSUE-{i}"
            parts = []
            if fields.get("summary"):
                parts.append(f"Summary: {fields['summary']}")
            for label, value in (
                ("Description", fields.get("description")),
                ("Acceptance Criteria", fields.get("acceptance_criteria")
                    or fields.get("acceptanceCriteria")),
                ("Priority", (fields.get("priority") or {}).get("name")
                    if isinstance(fields.get("priority"), dict) else fields.get("priority")),
                ("Component", ", ".join(
                    c.get("name", "") if isinstance(c, dict) else str(c)
                    for c in (fields.get("components") or [])) or None),
            ):
                if value:
                    parts.append(f"{label}: {value}")
            components = fields.get("components") or []
            hint = None
            if components:
                first = components[0]
                hint = first.get("name") if isinstance(first, dict) else str(first)
            body = "\n".join(p for p in parts if p.strip())
            if not body.strip():
                doc.warnings.append(f"{key}: issue has no summary, description or acceptance criteria")
                continue
            doc.blocks.append(Block(ref=str(key), text=body, hint=hint))
        return doc


_REGISTRY: Dict[str, IDocumentParser] = {}


def register(parser: IDocumentParser) -> None:
    for ext in parser.extensions:
        _REGISTRY[ext] = parser


for _p in (TextParser(), DocxParser(), PdfParser(), XlsxParser(), JiraJsonParser()):
    register(_p)


def supported_extensions() -> List[str]:
    return sorted(_REGISTRY)


def parse_document(path: str) -> ParsedDocument:
    """Dispatch on file extension. Raises for unsupported formats."""
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    ext = os.path.splitext(path)[1].lower()
    parser = _REGISTRY.get(ext)
    if parser is None:
        raise ValueError(
            f"No parser registered for '{ext}'. Supported: {', '.join(supported_extensions())}"
        )
    return parser.parse(path)
