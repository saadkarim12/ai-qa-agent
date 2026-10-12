"""Test case generator tests.

The model itself is replaced by a stand-in that returns fixed drafts, so these
tests check everything around it: story parsing, the grounding checks, numbering,
the markdown hand-off to Playwright, and failure handling.
"""

import os
import sys
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tcgen.examples import example_text_for_site, exemplars
from tcgen.guard import check_set, grounding_issue_count
from tcgen.llm import ModelError
from tcgen.markdown import COLUMNS, parse_markdown, render_markdown
from tcgen.models import DraftCase, DraftSet, TestCase, TestDatum
from tcgen.pipeline import finalise, generate
from tcgen.prompt import build_messages
from tcgen.stories import Story, parse_stories

SITE = "https://www.saucedemo.com/"
USER_FORMAT = """Site URL: https://www.saucedemo.com/
User Story 1 — User Login
As a registered user
I want to log in to the website using valid credentials
So that I can access the application.

As a logged-in user
I want to add a product to my shopping cart
So that I can purchase the selected product."""


def _draft(**over):
    base = dict(
        objective="Log in with valid credentials", derivation="stated", category="positive",
        preconditions=["User is not logged in"],
        steps=["Navigate to %s" % SITE, "Enter `standard_user` in the Username field",
               "Enter `secret_sauce` in the Password field", "Click the Login button",
               "Verify the Products page is displayed"],
        test_data=[TestDatum(name="username", value="standard_user"),
                   TestDatum(name="password", value="secret_sauce")],
        expected_result="User is redirected to the Products page",
        risk="Critical", priority="P1", evidence_requirements=["Screenshot"],
    )
    base.update(over)
    return DraftCase(**base)


def _case(story_id="US-1", **over):
    d = _draft(**over)
    return TestCase(test_id="TC-X-001", story_id=story_id, source="", **d.model_dump())


class _FakeStream:
    """Mimics the SDK's stream context manager: JSON arrives in small chunks."""

    def __init__(self, text, model):
        self._text, self._model = text, model

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    @property
    def text_stream(self):
        for i in range(0, len(self._text), 40):
            yield self._text[i:i + 40]

    def get_final_message(self):
        return SimpleNamespace(
            stop_reason="end_turn", model=self._model,
            usage=SimpleNamespace(input_tokens=100, output_tokens=50,
                                  cache_creation_input_tokens=0, cache_read_input_tokens=0))


class FakeClient:
    """Stands in for anthropic.Anthropic.

    drafts[n-1] answers "Story US-n", so results don't depend on which of the
    parallel requests happens to be sent first.
    """

    def __init__(self, drafts, fail_on=None):
        self._drafts = list(drafts)
        self._fail_on = fail_on or set()
        self.calls = []
        self.warmed = []
        self.beta = SimpleNamespace(messages=self)
        self.messages = self

    def stream(self, **kwargs):
        self.calls.append(kwargs)
        story_line = kwargs["messages"][-1]["content"]
        if any(tag in story_line for tag in self._fail_on):
            raise ModelError("simulated failure")
        n = int(story_line.split("Story US-")[1].split(":")[0])
        return _FakeStream(self._drafts[n - 1].model_dump_json(), kwargs["model"])

    def create(self, **kwargs):
        self.warmed.append(kwargs)
        return SimpleNamespace(stop_reason="max_tokens")


# --- stories -----------------------------------------------------------------

def test_parses_the_users_own_format():
    stories = parse_stories(USER_FORMAT)
    assert [s.id for s in stories] == ["US-1", "US-2"]
    assert stories[0].title == "User Login"
    assert all(s.site_url == SITE for s in stories)
    assert "add a product" in stories[1].text


def test_site_field_is_used_when_the_text_has_none():
    stories = parse_stories("As a user\nI want to search\nSo that I find things.",
                            site_url="https://shop.example/")
    assert stories[0].site_url == "https://shop.example/"


def test_empty_input_gives_no_stories():
    assert parse_stories("   \n\n") == []


# --- examples ----------------------------------------------------------------

def test_exemplars_match_the_schema_and_the_users_counts():
    ex = exemplars()
    assert [len(e["cases"]) for e in ex] == [8, 5]


def test_example_data_never_applies_to_another_site():
    assert example_text_for_site("https://shop.example/") == ""


def test_held_out_answers_are_not_in_the_grounding_text():
    only_cart = example_text_for_site(SITE, [exemplars()[1]])
    assert "locked_out_user" not in only_cart
    assert "Epic sadface" not in only_cart


# --- prompt ------------------------------------------------------------------

def test_examples_are_cached_and_the_story_comes_last():
    msgs = build_messages(Story("US-1", "t", "As a user\nI want X", SITE), exemplars())
    assert [m["role"] for m in msgs] == ["user", "assistant", "user", "assistant", "user"]
    assert "cache_control" in msgs[-2]["content"][0]
    assert "Story US-1" in msgs[-1]["content"]


def test_a_different_site_is_called_out_to_the_model():
    msgs = build_messages(Story("US-1", "t", "As a user\nI want X", "https://shop.example/"),
                          exemplars())
    assert "different site" in msgs[-1]["content"]


# --- grounding checks --------------------------------------------------------

def _issues(case, corpus):
    check_set([case], corpus)
    return case.issues


def test_clean_case_has_no_issues():
    corpus = USER_FORMAT + example_text_for_site(SITE)
    assert _issues(_case(), corpus) == []


def test_invented_test_data_is_flagged():
    case = _case(test_data=[TestDatum(name="username", value="problem_user")])
    issues = _issues(case, USER_FORMAT)
    assert any("problem_user" in i and i.startswith("grounding") for i in issues)


def test_placeholders_are_never_flagged():
    case = _case(steps=["Enter `<valid username>` in the Username field"],
                 test_data=[TestDatum(name="username", value="<valid username>")])
    assert not [i for i in _issues(case, "") if i.startswith("grounding")]


def test_quoted_on_screen_text_must_come_from_the_input():
    invented = _case(expected_result='The message "Sorry, try again later" is shown')
    assert any("Sorry, try again later" in i for i in _issues(invented, USER_FORMAT))

    observed = _case(expected_result='The message "Epic sadface: Username is required" appears')
    corpus = example_text_for_site(SITE)
    assert not any("Epic sadface" in i for i in _issues(observed, corpus))


def test_urls_must_come_from_the_input():
    case = _case(steps=["Navigate to https://www.saucedemo.com/admin.html"])
    assert any("admin.html" in i for i in _issues(case, USER_FORMAT))


def test_example_accounts_are_flagged_on_another_site():
    case = _case(steps=["Navigate to https://shop.example/"])
    story = "Site URL: https://shop.example/\nAs a user\nI want to log in"
    corpus = story + example_text_for_site("https://shop.example/")
    issues = _issues(case, corpus)
    assert any("standard_user" in i for i in issues)


def test_steps_must_start_with_an_action_verb():
    case = _case(steps=["The user logs in somehow"])
    assert any(i.startswith("format") and "step 1" in i for i in _issues(case, USER_FORMAT))


def test_repeated_objective_is_flagged():
    a, b = _case(), _case()
    check_set([a, b], USER_FORMAT + example_text_for_site(SITE))
    assert not a.issues and any("repeats" in i for i in b.issues)


# --- markdown hand-off -------------------------------------------------------

def test_markdown_has_exactly_the_agreed_columns():
    story = parse_stories(USER_FORMAT)[0]
    md = render_markdown([story], [_case()])
    header = [l for l in md.splitlines() if l.startswith("| Test ID")][0]
    assert [c.strip() for c in header.strip("|").split("|")] == COLUMNS


def test_every_field_survives_a_round_trip_through_markdown():
    story = parse_stories(USER_FORMAT)[0]
    tricky = _case(
        objective="Reject an SQL | injection-style username",
        test_data=[TestDatum(name="username", value="' OR '1'='1"),
                   TestDatum(name="password", value="<valid password>")],
        preconditions=["A", "B"], derivation="derived", category="security",
    )
    tricky.source = "US-1 · derived (security)"
    back = parse_markdown(render_markdown([story], [tricky]))
    assert len(back) == 1
    got = back[0]
    for f in ("test_id", "story_id", "objective", "preconditions", "steps", "test_data",
              "expected_result", "risk", "priority", "evidence_requirements", "status",
              "category", "derivation"):
        assert getattr(got, f) == getattr(tricky, f), f


def test_flagged_cases_get_review_notes_outside_the_table():
    story = parse_stories(USER_FORMAT)[0]
    case = _case(test_data=[TestDatum(name="username", value="problem_user")])
    check_set([case], USER_FORMAT)
    md = render_markdown([story], [case])
    assert "Check before running" in md
    assert len(parse_markdown(md)) == 1  # notes don't break the table


# --- pipeline ----------------------------------------------------------------

def test_generate_numbers_labels_and_checks():
    client = FakeClient([
        DraftSet(area_code="login", cases=[_draft(), _draft(objective="Reject wrong password",
                                                           derivation="derived",
                                                           category="negative")]),
        DraftSet(area_code="CART", cases=[_draft(objective="Add one product")]),
    ])
    result = generate(USER_FORMAT, None, client, model_key="sonnet")
    assert [c.test_id for c in result.cases] == ["TC-LOGIN-001", "TC-LOGIN-002", "TC-CART-001"]
    assert result.cases[1].source == "US-1 · derived (negative)"
    assert all(c.status == "Draft" for c in result.cases)
    assert result.usage.input_tokens == 200
    assert len(parse_markdown(result.markdown)) == 3


def test_one_failed_story_does_not_lose_the_others():
    client = FakeClient([DraftSet(area_code="LOGIN", cases=[_draft()])],
                        fail_on={"Story US-2"})
    result = generate(USER_FORMAT, None, client, model_key="sonnet")
    assert "US-2" in result.errors
    assert [c.story_id for c in result.cases] == ["US-1"]


def test_missing_key_is_explained():
    with pytest.raises(ModelError, match="ANTHROPIC_API_KEY"):
        generate(USER_FORMAT, None, None)


def test_no_story_is_explained():
    with pytest.raises(ModelError, match="No user story"):
        generate("   ", None, FakeClient([]))


def test_grounding_issue_count():
    case = _case(test_data=[TestDatum(name="username", value="problem_user")])
    check_set([case], USER_FORMAT)
    assert grounding_issue_count([case]) >= 1


# --- server ------------------------------------------------------------------

import base64
import http.server
import json as _json
import socket
import threading
import urllib.error
import urllib.request

import tcgen.server as server_mod


@pytest.fixture
def served(monkeypatch):
    """Run the real server on a free port with the stand-in model."""
    drafts = [DraftSet(area_code="LOGIN", cases=[_draft()]),
              DraftSet(area_code="CART", cases=[_draft(objective="Add one product")])]
    monkeypatch.setattr(server_mod, "make_client", lambda: FakeClient(list(drafts)))
    monkeypatch.setattr(server_mod, "_spend", {"day": None, "usd": 0.0})
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", port), server_mod.Handler)
    t = threading.Thread(target=httpd.serve_forever, daemon=True)
    t.start()
    yield "http://127.0.0.1:%d" % port
    httpd.shutdown()
    httpd.server_close()


def _req(url, body=None, password=None):
    data = _json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data,
                                 headers={"Content-Type": "application/json"})
    if password is not None:
        token = base64.b64encode(("qa:%s" % password).encode()).decode()
        req.add_header("Authorization", "Basic " + token)
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def test_page_and_config_are_served(served):
    status, body = _req(served + "/")
    assert status == 200 and b"Test Case Writer" in body
    status, body = _req(served + "/api/config")
    cfg = _json.loads(body)
    assert cfg["configured"] is True
    assert "Site URL:" in cfg["example_stories"]


def test_generate_returns_cases_and_markdown(served):
    status, body = _req(served + "/api/generate", {"stories": USER_FORMAT})
    d = _json.loads(body)
    assert status == 200
    assert [c["test_id"] for c in d["cases"]] == ["TC-LOGIN-001", "TC-CART-001"]
    assert "| Test ID |" in d["markdown"]
    assert d["usage"]["input_tokens"] == 200


def test_password_is_enforced_when_set(served, monkeypatch):
    monkeypatch.setenv("TCGEN_PASSWORD", "s3cret")
    assert _req(served + "/")[0] == 401
    assert _req(served + "/", password="wrong")[0] == 401
    assert _req(served + "/", password="s3cret")[0] == 200
    assert _req(served + "/healthz")[0] == 200  # hosts need this without a password


def test_daily_spending_cap_stops_generation(served, monkeypatch):
    monkeypatch.setenv("TCGEN_DAILY_USD", "0")
    status, body = _req(served + "/api/generate", {"stories": USER_FORMAT})
    assert status == 429
    assert "spending limit" in _json.loads(body)["error"]


def test_empty_stories_get_a_plain_explanation(served):
    status, body = _req(served + "/api/generate", {"stories": "  "})
    assert status == 400
    assert "No user story" in _json.loads(body)["error"]


def test_missing_key_is_reported_by_config(served, monkeypatch):
    monkeypatch.setattr(server_mod, "make_client", lambda: None)
    cfg = _json.loads(_req(served + "/api/config")[1])
    assert cfg["configured"] is False


# --- checker refinements found by reading real model output ------------------

def _grounding(case, corpus, site=SITE):
    check_set([case], corpus, site)
    return [i for i in case.issues if i.startswith("grounding")]


def test_generic_attack_strings_are_not_hallucinations():
    for payload in ["' OR '1'='1", "<script>alert(1)</script>", "../../etc/passwd"]:
        case = _case(test_data=[TestDatum(name="username", value=payload)],
                     steps=["Enter `%s` in the Username field" % payload])
        assert _grounding(case, USER_FORMAT) == [], payload


def test_case_variants_of_example_data_are_grounded():
    corpus = example_text_for_site(SITE)
    case = _case(test_data=[TestDatum(name="username", value="STANDARD_USER")])
    assert _grounding(case, corpus) == []


def test_site_url_plus_a_path_from_the_examples_is_grounded():
    corpus = example_text_for_site(SITE)  # the examples mention '/cart.html'
    case = _case(steps=["Navigate to https://www.saucedemo.com/cart.html"])
    assert _grounding(case, corpus) == []


def test_a_path_on_the_site_that_nobody_mentioned_is_still_flagged():
    case = _case(steps=["Navigate to https://www.saucedemo.com/admin.html"])
    assert any("admin.html" in i for i in _grounding(case, example_text_for_site(SITE)))


def test_product_names_from_memory_are_flagged():
    case = _case(test_data=[TestDatum(name="product", value="Sauce Labs Backpack")])
    assert any("Sauce Labs Backpack" in i for i in _grounding(case, example_text_for_site(SITE)))


def test_one_invention_reused_across_cases_counts_once():
    a = _case(objective="A", test_data=[TestDatum(name="username", value="problem_user")])
    b = _case(objective="B", test_data=[TestDatum(name="username", value="problem_user")])
    check_set([a, b], USER_FORMAT + example_text_for_site(SITE), SITE)
    assert grounding_issue_count([a, b]) == 1  # not 2


def test_prompt_warns_against_outside_knowledge_of_the_site():
    from tcgen.prompt import SYSTEM
    assert "You may recognise the website" in SYSTEM


def test_keyboard_keys_are_not_hallucinations():
    case = _case(steps=["Enter `standard_user` in the Username field", "Press `Enter`"])
    assert _grounding(case, example_text_for_site(SITE)) == []


def test_padded_variant_of_example_data_is_grounded():
    case = _case(test_data=[TestDatum(name="username", value="` standard_user `")])
    assert _grounding(case, example_text_for_site(SITE)) == []


def test_opus_requests_fallback_and_less_thinking():
    client = FakeClient([DraftSet(area_code="LOGIN", cases=[_draft()])])
    generate(USER_FORMAT.split("\n\nAs a logged-in")[0], None, client, model_key="opus")
    sent = client.calls[0]
    assert sent["betas"] == ["server-side-fallback-2026-07-01"]
    assert sent["extra_body"] == {"fallbacks": "default"}
    assert sent["output_config"]["effort"] == "low"
    assert sent["output_config"]["format"]["type"] == "json_schema"


def test_cheaper_models_send_no_fallback():
    client = FakeClient([DraftSet(area_code="LOGIN", cases=[_draft()])])
    generate(USER_FORMAT.split("\n\nAs a logged-in")[0], None, client, model_key="haiku")
    assert "betas" not in client.calls[0]


# --- streaming ----------------------------------------------------------------

from tcgen.llm import CaseStreamParser


def test_stream_parser_recovers_every_case_from_small_chunks():
    full = DraftSet(area_code="LOGIN", cases=[DraftCase(**c) for c in exemplars()[0]["cases"]])
    text = full.model_dump_json()
    parser, got = CaseStreamParser(), []
    for i in range(0, len(text), 7):
        got += parser.feed(text[i:i + 7])
    assert [c.objective for c in got] == [c.objective for c in full.cases]


def test_stream_parser_is_not_fooled_by_braces_in_text():
    tricky = DraftSet(area_code="X", cases=[_draft(objective='A } brace, a " quote, a {b}')])
    parser, got = CaseStreamParser(), []
    for ch in tricky.model_dump_json():
        got += parser.feed(ch)
    assert len(got) == 1 and got[0].objective == 'A } brace, a " quote, a {b}'


def test_events_arrive_in_a_usable_order():
    client = FakeClient([DraftSet(area_code="LOGIN", cases=[_draft(), _draft(objective="B")]),
                         DraftSet(area_code="CART", cases=[_draft(objective="C")])])
    events = []
    generate(USER_FORMAT, None, client, model_key="haiku", on_event=events.append)
    kinds = [e["type"] for e in events]
    assert kinds[0] == "stories"
    assert kinds.count("case") == 3 and kinds.count("story_done") == 2
    # each story's cases are reported before that story is marked done
    for sid in ("US-1", "US-2"):
        idx = [i for i, e in enumerate(events) if e.get("story_id") == sid]
        assert events[idx[-1]]["type"] == "story_done"


def test_live_cases_are_checked_as_they_arrive():
    bad = _draft(test_data=[TestDatum(name="username", value="problem_user")])
    client = FakeClient([DraftSet(area_code="LOGIN", cases=[bad])])
    events = []
    generate(USER_FORMAT.split("\n\nAs a logged-in")[0], None, client, model_key="haiku",
             on_event=events.append)
    live = [e for e in events if e["type"] == "case"][0]["case"]
    assert any("problem_user" in i for i in live["issues"])


def test_final_ids_follow_story_order_not_finish_order():
    client = FakeClient([DraftSet(area_code="LOGIN", cases=[_draft()]),
                         DraftSet(area_code="LOGIN", cases=[_draft(objective="Other")])])
    result = generate(USER_FORMAT, None, client, model_key="haiku")
    assert [(c.story_id, c.test_id) for c in result.cases] == [
        ("US-1", "TC-LOGIN-001"), ("US-2", "TC-LOGIN-002")]


# --- upload -------------------------------------------------------------------

from tcgen.extract import extract_text

FEWSHOT_XLSX = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "sample_data", "fewshot", "Sample-US-TC.xlsx")


def test_workbook_upload_keeps_stories_and_skips_test_case_sheets():
    text, notes = extract_text("Sample-US-TC.xlsx", open(FEWSHOT_XLSX, "rb").read())
    assert len(parse_stories(text)) == 2
    assert "TC-LOGIN" not in text
    assert any("TC1" in n and "TC2" in n for n in notes)


def test_word_upload(tmp_path):
    import docx
    d = docx.Document()
    for line in ["Site URL: https://shop.example/", "User Story 3 — Checkout",
                 "As a customer", "I want to pay by card", "So that I get my order."]:
        d.add_paragraph(line)
    path = tmp_path / "story.docx"
    d.save(str(path))
    text, _ = extract_text("story.docx", path.read_bytes())
    stories = parse_stories(text)
    assert stories[0].title == "Checkout" and stories[0].site_url == "https://shop.example/"


def test_plain_text_upload():
    text, _ = extract_text("s.txt", b"As a user\nI want to search\nSo that I find things.")
    assert len(parse_stories(text)) == 1


def test_unsupported_and_empty_uploads_are_explained():
    with pytest.raises(ValueError, match="isn't supported"):
        extract_text("s.exe", b"x")
    with pytest.raises(ValueError, match="No text found"):
        extract_text("s.txt", b"   ")



def _sse(url, body):
    req = urllib.request.Request(url, data=_json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=15) as r:
        raw = r.read().decode()
    return [_json.loads(l[6:]) for l in raw.split("\n") if l.startswith("data: ")]


def test_streaming_endpoint_sends_cases_then_the_result(served):
    events = _sse(served + "/api/generate/stream", {"stories": USER_FORMAT, "model": "haiku"})
    kinds = [e["type"] for e in events]
    assert kinds[0] == "stories" and kinds[-1] == "done"
    assert kinds.count("case") == 2
    assert [c["test_id"] for c in events[-1]["cases"]] == ["TC-LOGIN-001", "TC-CART-001"]


def test_config_lists_the_models_to_choose_from(served):
    cfg = _json.loads(_req(served + "/api/config")[1])
    assert [m["key"] for m in cfg["models"]] == ["opus", "sonnet", "haiku"]
    assert cfg["default_model"] == "opus"
    assert ".docx" in cfg["upload_types"]


def test_unknown_model_is_refused(served):
    status, body = _req(served + "/api/generate", {"stories": USER_FORMAT, "model": "gpt-4"})
    assert status == 400 and "Choose one of" in _json.loads(body)["error"]


def test_warm_endpoint_accepts_a_model(served):
    status, body = _req(served + "/api/warm", {"model": "haiku"})
    assert status == 202


def test_extract_endpoint_returns_story_text(served):
    b64 = base64.b64encode(open(FEWSHOT_XLSX, "rb").read()).decode()
    status, body = _req(served + "/api/extract", {"filename": "US.xlsx", "content_b64": b64})
    d = _json.loads(body)
    assert status == 200 and d["stories_found"] == 2 and d["notes"]


def test_boundary_numbers_are_not_hallucinations():
    case = _case(steps=["Enter `-1` in the quantity field", "Enter `101` in the quantity field"],
                 test_data=[TestDatum(name="quantity", value="0")])
    assert _grounding(case, example_text_for_site(SITE)) == []


def test_non_standard_step_verbs_are_flagged_as_format():
    case = _case(steps=["Attempt to add a product", "Note the cart count"])
    fmt = [i for i in _issues(case, example_text_for_site(SITE)) if i.startswith("format")]
    assert len(fmt) == 2


def test_hosted_without_a_password_refuses_everything(served, monkeypatch):
    monkeypatch.setenv("PORT", "10000")
    monkeypatch.delenv("TCGEN_PASSWORD", raising=False)
    assert _req(served + "/")[0] == 503
    assert _req(served + "/api/generate", {"stories": USER_FORMAT})[0] == 503
    assert _req(served + "/healthz")[0] == 200  # the host's health check still works


def test_hosted_with_a_password_works(served, monkeypatch):
    monkeypatch.setenv("PORT", "10000")
    monkeypatch.setenv("TCGEN_PASSWORD", "team-pass")
    assert _req(served + "/")[0] == 401
    assert _req(served + "/", password="team-pass")[0] == 200
